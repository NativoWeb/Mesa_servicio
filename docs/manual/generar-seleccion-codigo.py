"""
Genera la Selección documental del código fuente de Mesa de Servicio TI (.docx).
Replica visualmente el formato de NW GCA REG 03 (Selección GCA).
NATIVOWEB S.A.S. — NW MDA REG 03
"""

import os
import hashlib
from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import nsdecls
from docx.oxml import parse_xml

BASE = os.path.dirname(os.path.abspath(__file__))
PROJECT = os.path.dirname(os.path.dirname(BASE))  # mesa-de-ayuda-uts
ASSETS = os.path.join(BASE, "assets")
LOGO_NATIVO = r"D:\MARLY\DOCUMENTOS PROYECTOS\NATIVOWEB\LOGO NATIVO.png"
LOGO_HEADER = os.path.join(ASSETS, "logo_header.png")
LOGO_COVER = os.path.join(ASSETS, "logo_cover.png")
OUTPUT = os.path.join(BASE, "NW_MDA_REG_03_Seleccion_Codigo_Fuente_Mesa_Servicio_TI_1_0_0.docx")

if not os.path.exists(LOGO_COVER) and os.path.exists(LOGO_NATIVO):
    LOGO_COVER = LOGO_NATIVO
if not os.path.exists(LOGO_HEADER) and os.path.exists(LOGO_NATIVO):
    LOGO_HEADER = LOGO_NATIVO

# Colores
TABLE_HEADER_BG = "4472C4"
BLACK = RGBColor(0x00, 0x00, 0x00)
GRAY_DARK = RGBColor(0x33, 0x33, 0x33)
GRAY_MED = RGBColor(0x66, 0x66, 0x66)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CODE_BG = "F2F2F2"

doc = Document()

# ── Estilos base ─────────────────────────────────────────
style_normal = doc.styles['Normal']
style_normal.font.name = 'Calibri'
style_normal.font.size = Pt(10)
style_normal.font.color.rgb = GRAY_DARK
style_normal.paragraph_format.space_after = Pt(6)
style_normal.paragraph_format.space_before = Pt(0)
style_normal.paragraph_format.line_spacing = 1.15

for level, size in [(1, 18), (2, 13), (3, 11)]:
    h = doc.styles[f'Heading {level}']
    h.font.name = 'Calibri'
    h.font.size = Pt(size)
    h.font.color.rgb = BLACK
    h.font.bold = True
    h.paragraph_format.space_before = Pt(14 if level == 1 else 10)
    h.paragraph_format.space_after = Pt(4)

for section in doc.sections:
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.54)
    section.right_margin = Cm(2.54)
    section.header_distance = Cm(1.0)
    section.footer_distance = Cm(1.0)

# ── Header ───────────────────────────────────────────────
header = doc.sections[0].header
header.is_linked_to_previous = False
ht = header.add_table(1, 2, width=Inches(6.5))
ht.alignment = WD_TABLE_ALIGNMENT.CENTER
c0 = ht.cell(0, 0)
c0.width = Inches(1.5)
p0 = c0.paragraphs[0]
p0.alignment = WD_ALIGN_PARAGRAPH.LEFT
if os.path.exists(LOGO_HEADER):
    r0 = p0.add_run()
    r0.add_picture(LOGO_HEADER, height=Inches(0.4))
c1 = ht.cell(0, 1)
p1 = c1.paragraphs[0]
p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r1 = p1.add_run("NATIVOWEB S.A.S.  |  NIT 901986453-2")
r1.font.size = Pt(8)
r1.font.color.rgb = GRAY_MED
for cell in ht.rows[0].cells:
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}>'
        '<w:top w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
        '<w:left w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
        '<w:bottom w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
        '<w:right w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
        '</w:tcBorders>')
    tcPr.append(tcBorders)

# ── Footer ───────────────────────────────────────────────
footer = doc.sections[0].footer
footer.is_linked_to_previous = False
ft = footer.add_table(1, 2, width=Inches(6.5))
ft.alignment = WD_TABLE_ALIGNMENT.CENTER
fc0 = ft.cell(0, 0)
fp0 = fc0.paragraphs[0]
fp0.alignment = WD_ALIGN_PARAGRAPH.LEFT
fr0 = fp0.add_run("NW MDA REG 03  ·  Selección de código")
fr0.font.size = Pt(8)
fr0.font.color.rgb = GRAY_MED
fc1 = ft.cell(0, 1)
fp1 = fc1.paragraphs[0]
fp1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
fr1a = fp1.add_run("Página ")
fr1a.font.size = Pt(8)
fr1a.font.color.rgb = GRAY_MED
fldChar1 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="begin"/>')
fr1b = fp1.add_run()
fr1b._r.append(fldChar1)
instrText = parse_xml(f'<w:instrText {nsdecls("w")} xml:space="preserve"> PAGE </w:instrText>')
fr1c = fp1.add_run()
fr1c._r.append(instrText)
fldChar2 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="end"/>')
fr1d = fp1.add_run()
fr1d._r.append(fldChar2)
fr1d.font.size = Pt(8)
for cell in ft.rows[0].cells:
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}>'
        '<w:top w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
        '<w:left w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
        '<w:bottom w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
        '<w:right w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
        '</w:tcBorders>')
    tcPr.append(tcBorders)


# ── Helpers ──────────────────────────────────────────────
def add_para(text, bold=False, size=10, color=GRAY_DARK, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.bold = bold
    return p

def set_cell_bg(cell, color_hex):
    shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}" w:val="clear"/>')
    cell._tc.get_or_add_tcPr().append(shading)

def add_styled_table(headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_bg(cell, TABLE_HEADER_BG)
        p = cell.paragraphs[0]
        run = p.add_run(h)
        run.font.bold = True
        run.font.size = Pt(9)
        run.font.color.rgb = WHITE
    for r, row_data in enumerate(rows):
        for c, val in enumerate(row_data):
            cell = table.rows[r + 1].cells[c]
            p = cell.paragraphs[0]
            run = p.add_run(val)
            run.font.size = Pt(9)
            run.font.color.rgb = GRAY_DARK
    tbl = table._tbl
    tblPr = tbl.tblPr if tbl.tblPr is not None else parse_xml(f'<w:tblPr {nsdecls("w")}/>')
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        '<w:top w:val="single" w:sz="4" w:space="0" w:color="BFBFBF"/>'
        '<w:left w:val="single" w:sz="4" w:space="0" w:color="BFBFBF"/>'
        '<w:bottom w:val="single" w:sz="4" w:space="0" w:color="BFBFBF"/>'
        '<w:right w:val="single" w:sz="4" w:space="0" w:color="BFBFBF"/>'
        '<w:insideH w:val="single" w:sz="4" w:space="0" w:color="BFBFBF"/>'
        '<w:insideV w:val="single" w:sz="4" w:space="0" w:color="BFBFBF"/>'
        '</w:tblBorders>'
    )
    tblPr.append(borders)
    return table

def page_break():
    doc.add_page_break()

def add_note(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    run = p.add_run(text)
    run.font.size = Pt(9)
    run.font.color.rgb = GRAY_MED
    run.font.italic = True


def read_source_file(relative_path):
    """Lee un archivo fuente y retorna contenido, líneas y hash SHA-256."""
    full_path = os.path.join(PROJECT, relative_path)
    with open(full_path, 'r', encoding='utf-8') as f:
        content = f.read()
    # Asegurar salto final de línea para el hash
    if not content.endswith('\n'):
        content_for_hash = content + '\n'
    else:
        content_for_hash = content
    sha256 = hashlib.sha256(content_for_hash.encode('utf-8')).hexdigest()
    lines = content.splitlines()
    return content, lines, len(lines), sha256


def add_code_block(lines, start, end, total):
    """Agrega un bloque de código con fondo gris y fuente monoespaciada."""
    for line in lines[start - 1 : end]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        # Fondo gris para el párrafo
        pPr = p._p.get_or_add_pPr()
        shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{CODE_BG}" w:val="clear"/>')
        pPr.append(shading)
        run = p.add_run(line if line else " ")
        run.font.name = 'Consolas'
        run.font.size = Pt(8)
        run.font.color.rgb = GRAY_DARK


def add_component(number, relative_path, lines, total_lines, sha256, max_lines_per_page=47):
    """
    Agrega un componente completo al documento, paginando si es necesario.
    Replica la estructura del GCA: título, ruta, rango de líneas, código, hash al final.
    """
    chunks = []
    start = 1
    while start <= total_lines:
        end = min(start + max_lines_per_page - 1, total_lines)
        chunks.append((start, end))
        start = end + 1

    for chunk_idx, (chunk_start, chunk_end) in enumerate(chunks):
        doc.add_heading(f"Componente {number}", level=1)
        add_para(relative_path, bold=True, size=11, color=BLACK,
                 align=WD_ALIGN_PARAGRAPH.LEFT, space_after=2)
        add_para(f"Líneas {chunk_start} a {chunk_end} de {total_lines}",
                 size=9, color=GRAY_MED, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=8)

        add_code_block(lines, chunk_start, chunk_end, total_lines)

        # Hash solo en la última página del componente
        if chunk_idx == len(chunks) - 1:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            run_label = p.add_run("SHA-256 de la transcripción")
            run_label.font.size = Pt(9)
            run_label.font.bold = True
            run_label.font.color.rgb = GRAY_DARK

            p_hash = doc.add_paragraph()
            p_hash.paragraph_format.space_after = Pt(4)
            run_hash = p_hash.add_run(sha256)
            run_hash.font.name = 'Consolas'
            run_hash.font.size = Pt(8)
            run_hash.font.color.rgb = GRAY_MED

        # Salto de página entre chunks y entre componentes
        if chunk_idx < len(chunks) - 1:
            page_break()

    page_break()


# ══════════════════════════════════════════════════════════
# DEFINICIÓN DE LOS 6 COMPONENTES
# ══════════════════════════════════════════════════════════

COMPONENTS = [
    {
        "path": "backend/app/Models/Ticket.php",
        "function": "Modelo principal de soporte técnico.",
    },
    {
        "path": "backend/app/Http/Controllers/Api/AuthController.php",
        "function": "Acceso, sesión y recuperación de contraseña.",
    },
    {
        "path": "backend/app/Services/SlaService.php",
        "function": "Cálculo de plazos y verificación de cumplimiento SLA.",
    },
    {
        "path": "backend/app/Enums/UserRole.php",
        "function": "Definición de roles y etiquetas del sistema.",
    },
    {
        "path": "frontend/src/lib/api.ts",
        "function": "Cliente HTTP con autenticación automática.",
    },
    {
        "path": "frontend/src/stores/auth-store.ts",
        "function": "Estado de sesión persistido y redirección por rol.",
    },
]

# Leer todos los archivos y calcular hashes
components_data = []
for comp in COMPONENTS:
    content, lines, total, sha256 = read_source_file(comp["path"])
    components_data.append({
        **comp,
        "content": content,
        "lines": lines,
        "total": total,
        "sha256": sha256,
    })

total_lines_all = sum(c["total"] for c in components_data)


# ══════════════════════════════════════════════════════════
# PORTADA
# ══════════════════════════════════════════════════════════

p_logo = doc.add_paragraph()
p_logo.alignment = WD_ALIGN_PARAGRAPH.LEFT
p_logo.paragraph_format.space_before = Pt(40)
if os.path.exists(LOGO_COVER):
    r_logo = p_logo.add_run()
    r_logo.add_picture(LOGO_COVER, width=Inches(2.5))

doc.add_paragraph()

p_title = doc.add_paragraph()
p_title.alignment = WD_ALIGN_PARAGRAPH.LEFT
run_title = p_title.add_run("Mesa de Servicio TI")
run_title.font.size = Pt(26)
run_title.font.bold = True
run_title.font.color.rgb = BLACK

p_sub = doc.add_paragraph()
p_sub.alignment = WD_ALIGN_PARAGRAPH.LEFT
run_sub = p_sub.add_run("Selección documental del código fuente")
run_sub.font.size = Pt(14)
run_sub.font.color.rgb = GRAY_DARK

p_ver = doc.add_paragraph()
run_ver = p_ver.add_run("Versión del software 1.0.0  |  Obra terminada en 2026")
run_ver.font.size = Pt(10)
run_ver.font.color.rgb = GRAY_DARK

add_para(
    "Selección de seis componentes del programa. Complementa la descripción técnica con la expresión escrita de "
    "funciones de soporte técnico, autenticación, nivel de servicio, definición de roles, acceso a datos "
    "y gestión de estado.",
    size=10
)

doc.add_paragraph()

add_styled_table(
    ["Identificación", "Detalle"],
    [
        ["Empresa", "NATIVOWEB S.A.S."],
        ["NIT", "901986453-2"],
        ["Autores", "BRAYAN ALFONSO APARICIO DURAN\nMARLY TATIANA RANGEL MORENO"],
        ["Edición documental", "Septiembre de 2026"],
        ["Código documental", "NW MDA REG 03"],
    ]
)

doc.add_paragraph()
add_para("Bucaramanga  ·  Colombia", size=10, color=GRAY_DARK)

page_break()


# ══════════════════════════════════════════════════════════
# ALCANCE E INVENTARIO
# ══════════════════════════════════════════════════════════
doc.add_heading("Alcance e inventario de la selección", level=1)

add_para(
    "Este documento reproduce seis bloques de código contenidos en la documentación técnica de la versión "
    "1.0.0. Constituye una selección parcial del programa y no un proyecto instalable completo. Las rutas "
    "identifican el lugar funcional del componente en la estructura del producto."
)

# Tabla de inventario
add_styled_table(
    ["Componente", "Función", "Líneas"],
    [
        [c["path"], c["function"], str(c["total"])]
        for c in components_data
    ]
)

doc.add_heading("Integridad de la reproducción", level=2)
add_para(
    "Los bloques conservan el texto de la selección aportada; no se completan archivos truncados ni se añaden "
    "funciones. Las huellas SHA-256 corresponden a la transcripción de cada bloque en UTF-8 con salto final de "
    "línea. No identifican un commit del repositorio ni una compilación del producto."
)
add_para(
    "Los derechos de Nativoweb se refieren a sus componentes propios. Las importaciones de librerías y "
    "dependencias no implican una atribución de autoría sobre software de terceros."
)

doc.add_heading("Ubicación de los originales transcritos", level=2)
add_para(
    "La carpeta CODIGO_SELECCION del expediente conserva los seis archivos de texto con sus rutas y un "
    "manifiesto de huellas. Puede emplearse para cotejar esta reproducción documental. El soporte principal "
    "recomendado para el registro del programa es su descripción técnica y funcional."
)

page_break()


# ══════════════════════════════════════════════════════════
# COMPONENTES DE CÓDIGO
# ══════════════════════════════════════════════════════════

for i, comp in enumerate(components_data, 1):
    add_component(i, comp["path"], comp["lines"], comp["total"], comp["sha256"])


# ── Guardar ──────────────────────────────────────────────
doc.save(OUTPUT)
print(f"Documento generado: {OUTPUT}")
print(f"Tamaño: {os.path.getsize(OUTPUT) / 1024 / 1024:.1f} MB")
print(f"\nResumen de componentes ({total_lines_all} líneas totales):")
for i, c in enumerate(components_data, 1):
    print(f"  {i}. {c['path']} ({c['total']} líneas)")
    print(f"     SHA-256: {c['sha256']}")
